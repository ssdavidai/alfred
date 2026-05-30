"""Tests for FileExtractionWorkflow + file_extraction activities.

Covered surfaces
----------------

* Pure extractor dispatch (classify_mime + extract_text_blob):
  - utf-8 text round-trips identically.
  - PDF fixture is parsed (single-page, embedded text) by pypdf.
  - DOCX fixture is parsed by python-docx.
  - XLSX fixture is parsed by openpyxl.
  - Unsupported mime raises ExtractionError("unsupported_mime").
  - PDF with no text layer raises ExtractionError("empty_pdf").
  - Binary garbage on a text/* mime is decoded with replacement (no
    crash); on a pdf mime it raises extractor_failed.

* Activity wrappers via ``ActivityEnvironment`` with httpx stubbed:
  - read_file_metadata locates the row in the /list response.
  - stamp_extraction_result PATCHes the right body shape.

* Workflow via WorkflowEnvironment.start_time_skipping() with stub
  activities:
  - happy path: success stamp with alfred_read_at + summary, no error.
  - unsupported mime short-circuits at fetch_and_extract_text.
  - empty extraction stamps the error code.
  - summariser failure stamps "summariser_failed".
  - stamp failure still returns a result (no exception).
"""
from __future__ import annotations

import asyncio
import io
import uuid
from typing import Any

import httpx
import pytest
from temporalio import activity
from temporalio.client import Client
from temporalio.testing import ActivityEnvironment, WorkflowEnvironment
from temporalio.worker import Worker

from src.activities.file_extraction import (
    ExtractionError,
    classify_mime,
    extract_text_blob,
    fetch_and_extract_text,
    read_file_metadata,
    stamp_extraction_result,
)
from src.workflows.file_extraction import (
    FileExtractionResult,
    FileExtractionWorkflow,
)


# ---------------------------------------------------------------------------
# Fixture builders — keep them in-test so the fixtures dir doesn't bloat.
# ---------------------------------------------------------------------------


def _build_min_pdf_with_text(text: str) -> bytes:
    """Build the minimum possible PDF that pypdf will extract `text` from.

    Hand-rolled (~600 bytes) so the test stays hermetic — no fixture
    file, no external generator. The structure follows ISO 32000-1's
    "minimum file" example: 1 page, 1 font, a single Tj operator.
    """
    objs: list[bytes] = []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>")
    objs.append(
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 200 200] /Contents 5 0 R >>"
    )
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = (
        b"BT /F1 12 Tf 50 100 Td (" + safe.encode("ascii", errors="replace") + b") Tj ET"
    )
    obj5 = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    objs.append(obj5)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")
    xref_off = out.tell()
    out.write(b"xref\n")
    out.write(f"0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(b"trailer\n")
    out.write(f"<< /Size {len(objs) + 1} /Root 1 0 R >>\n".encode())
    out.write(b"startxref\n")
    out.write(f"{xref_off}\n".encode())
    out.write(b"%%EOF\n")
    return out.getvalue()


def _build_min_pdf_no_text_layer() -> bytes:
    """An "image-only" PDF with one page and zero text operators —
    pypdf returns '' for every page."""
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R >>",
        b"<< /Length 0 >>\nstream\n\nendstream",
    ]
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")
    xref_off = out.tell()
    out.write(b"xref\n")
    out.write(f"0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(b"trailer\n")
    out.write(f"<< /Size {len(objs) + 1} /Root 1 0 R >>\n".encode())
    out.write(b"startxref\n")
    out.write(f"{xref_off}\n".encode())
    out.write(b"%%EOF\n")
    return out.getvalue()


def _build_min_docx(text: str) -> bytes:
    """python-docx writes via the ``docx.Document`` API; use that
    directly to spit out a minimum DOCX containing a single paragraph.
    """
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_min_xlsx(values: list[list[Any]]) -> bytes:
    """openpyxl helper — single-sheet workbook with the given rows."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in values:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1. Pure extractor dispatch — no temporal, no httpx.
# ---------------------------------------------------------------------------


class TestClassifyMime:
    @pytest.mark.parametrize(
        "ct,expected",
        [
            ("text/plain", "text"),
            ("text/markdown; charset=utf-8", "text"),
            ("application/json", "text"),
            ("application/xml", "text"),
            ("application/yaml", "text"),
            ("application/pdf", "pdf"),
            (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "docx",
            ),
            (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            ),
            ("image/png", "unsupported"),
            ("audio/mpeg", "unsupported"),
            ("", "unsupported"),
            (None, "unsupported"),
        ],
    )
    def test_dispatch(self, ct, expected):
        assert classify_mime(ct) == expected


class TestExtractText:
    def test_utf8_round_trips(self):
        body = "Hello Alfred — the audit probe at 03:42.\nLine 2."
        out = extract_text_blob("text/plain", body.encode("utf-8"))
        assert out == body

    def test_text_with_garbage_uses_replacement(self):
        body = b"Hello \xff World"
        out = extract_text_blob("text/plain", body)
        assert "Hello" in out and "World" in out


class TestExtractPdf:
    def test_extracts_embedded_text(self):
        pdf = _build_min_pdf_with_text("Quarterly contract")
        out = extract_text_blob("application/pdf", pdf)
        assert "Quarterly contract" in out

    def test_empty_pdf_raises_empty_pdf(self):
        pdf = _build_min_pdf_no_text_layer()
        with pytest.raises(ExtractionError) as ei:
            extract_text_blob("application/pdf", pdf)
        assert ei.value.code == "empty_pdf"

    def test_garbage_bytes_raise_extractor_failed(self):
        with pytest.raises(ExtractionError) as ei:
            extract_text_blob("application/pdf", b"not a pdf at all")
        assert ei.value.code == "extractor_failed"


class TestExtractDocx:
    def test_extracts_paragraph(self):
        docx = _build_min_docx("Letterpress briefing for 2026-05-30")
        out = extract_text_blob(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx,
        )
        assert "Letterpress briefing" in out


class TestExtractXlsx:
    def test_extracts_cells(self):
        xlsx = _build_min_xlsx([["Q1", "Q2"], [100, 200]])
        out = extract_text_blob(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx,
        )
        assert "Q1" in out
        assert "Q2" in out
        assert "100" in out


class TestUnsupportedMime:
    def test_image_raises(self):
        with pytest.raises(ExtractionError) as ei:
            extract_text_blob("image/png", b"\x89PNG\r\n\x1a\n...")
        assert ei.value.code == "unsupported_mime"


# ---------------------------------------------------------------------------
# 2. Activity isolation — stub httpx.
# ---------------------------------------------------------------------------


class _StubTransport(httpx.AsyncBaseTransport):
    def __init__(self) -> None:
        self.queued: list[httpx.Response] = []
        self.requests: list[httpx.Request] = []

    def queue(
        self,
        status: int = 200,
        json_payload: dict[str, Any] | list[Any] | None = None,
        content: bytes | None = None,
    ) -> None:
        if content is not None:
            self.queued.append(httpx.Response(status, content=content))
        else:
            self.queued.append(httpx.Response(status, json=json_payload or {}))

    async def handle_async_request(self, request: httpx.Request) -> httpx.Response:  # type: ignore[override]
        self.requests.append(request)
        if not self.queued:
            return httpx.Response(500, json={"error": "no stub queued"})
        return self.queued.pop(0)


@pytest.fixture
def stub_transport(monkeypatch):
    transport = _StubTransport()
    real_init = httpx.AsyncClient.__init__

    def patched_init(self, *args, **kwargs):  # type: ignore[no-untyped-def]
        kwargs["transport"] = transport
        real_init(self, *args, **kwargs)

    monkeypatch.setattr(httpx.AsyncClient, "__init__", patched_init)
    monkeypatch.setenv("ALFRED_CTRL_URL", "http://stub-ctrl:3100")
    monkeypatch.setenv("AAS_API_KEY", "test-key")
    return transport


class TestReadFileMetadata:
    def test_locates_matching_row(self, stub_transport):
        stub_transport.queue(
            200,
            {
                "items": [
                    {"id": "f-other", "content_type": "image/png"},
                    {
                        "id": "f-target",
                        "content_type": "application/pdf",
                        "path": "01H/foo.pdf",
                        "original_filename": "foo.pdf",
                    },
                ],
            },
        )
        env = ActivityEnvironment()
        row = asyncio.run(env.run(read_file_metadata, "f-target"))
        assert row["id"] == "f-target"
        assert row["content_type"] == "application/pdf"

    def test_missing_id_raises(self, stub_transport):
        stub_transport.queue(200, {"items": []})
        env = ActivityEnvironment()
        with pytest.raises(ExtractionError) as ei:
            asyncio.run(env.run(read_file_metadata, "nope"))
        assert ei.value.code == "not_found"


class TestFetchAndExtractText:
    def test_fetches_blob_and_decodes(self, stub_transport):
        stub_transport.queue(200, content=b"Hello extracted body")
        env = ActivityEnvironment()
        out = asyncio.run(
            env.run(
                fetch_and_extract_text,
                "f-1",
                "text/plain",
                "01H/foo.txt",
            )
        )
        assert out["text"] == "Hello extracted body"
        assert out["extractor"] == "text"
        # The request encoded each path segment.
        req = stub_transport.requests[0]
        assert "01H/foo.txt" in str(req.url)

    def test_unsupported_mime_short_circuits(self, stub_transport):
        # No blob fetch should happen.
        env = ActivityEnvironment()
        with pytest.raises(ExtractionError) as ei:
            asyncio.run(
                env.run(
                    fetch_and_extract_text,
                    "f-1",
                    "image/png",
                    "01H/foo.png",
                )
            )
        assert ei.value.code == "unsupported_mime"
        assert stub_transport.requests == []


class TestStampExtractionResult:
    def test_patches_correct_body(self, stub_transport):
        stub_transport.queue(
            200,
            {
                "id": "f-1",
                "alfred_read_at": 12345,
                "summary": "ok",
                "extraction_error": None,
            },
        )
        env = ActivityEnvironment()
        resp = asyncio.run(
            env.run(
                stamp_extraction_result,
                "f-1",
                12345,
                "ok",
                None,
            )
        )
        assert resp["alfred_read_at"] == 12345
        req = stub_transport.requests[0]
        assert req.method == "PATCH"
        assert "/api/v1/files/f-1/extraction" in str(req.url)
        # Body carries all three keys verbatim.
        import json as _json
        body = _json.loads(req.content.decode())
        assert body == {
            "alfred_read_at": 12345,
            "summary": "ok",
            "extraction_error": None,
        }


# ---------------------------------------------------------------------------
# 3. Workflow tests — stub activities; no httpx call at all.
# ---------------------------------------------------------------------------


def _make_stubs(
    *,
    metadata_row: dict[str, Any] | None = None,
    extract_payload: dict[str, Any] | None = None,
    extract_raises: BaseException | None = None,
    summary: str | None = None,
    summary_raises: BaseException | None = None,
    stamp_raises: BaseException | None = None,
    stamp_recorder: list[dict[str, Any]] | None = None,
):
    @activity.defn(name="read_file_metadata")
    async def stub_read(file_id: str) -> dict[str, Any]:
        return metadata_row or {
            "id": file_id,
            "content_type": "text/plain",
            "path": "01H/foo.txt",
            "original_filename": "foo.txt",
        }

    @activity.defn(name="fetch_and_extract_text")
    async def stub_extract(
        file_id: str, content_type: str | None, path: str,
    ) -> dict[str, Any]:
        if extract_raises is not None:
            raise extract_raises
        return extract_payload or {
            "text": "Hello body",
            "char_count": 10,
            "extractor": "text",
            "truncated_for_summary": False,
        }

    @activity.defn(name="summarise_extracted_text")
    async def stub_summary(
        file_id: str,
        extracted_text: str,
        original_filename: str | None,
        content_type: str | None,
        truncated: bool,
    ) -> str:
        if summary_raises is not None:
            raise summary_raises
        return summary or "A short summary."

    @activity.defn(name="stamp_extraction_result")
    async def stub_stamp(
        file_id: str,
        alfred_read_at: int | None,
        summary_text: str | None,
        extraction_error: str | None,
    ) -> dict[str, Any]:
        if stamp_recorder is not None:
            stamp_recorder.append(
                {
                    "file_id": file_id,
                    "alfred_read_at": alfred_read_at,
                    "summary": summary_text,
                    "extraction_error": extraction_error,
                }
            )
        if stamp_raises is not None:
            raise stamp_raises
        return {
            "id": file_id,
            "alfred_read_at": alfred_read_at,
            "summary": summary_text,
            "extraction_error": extraction_error,
        }

    return [stub_read, stub_extract, stub_summary, stub_stamp]


async def _run_workflow(
    stubs, file_id: str = "f-1",
) -> FileExtractionResult:
    async with await WorkflowEnvironment.start_time_skipping() as env:
        client: Client = env.client
        tq = f"file-extract-test-{uuid.uuid4()}"
        worker = Worker(
            client,
            task_queue=tq,
            workflows=[FileExtractionWorkflow],
            activities=stubs,
        )
        async with worker:
            return await client.execute_workflow(
                FileExtractionWorkflow.run,
                {"file_id": file_id},
                id=f"file-extract-run-{uuid.uuid4()}",
                task_queue=tq,
            )


def _to_dict(result: Any) -> dict[str, Any]:
    """Temporal can return either a FileExtractionResult or a dict
    depending on Python version + dataclass support; this normalises."""
    if isinstance(result, dict):
        return result
    return result.__dict__


class TestHappyPath:
    def test_stamps_alfred_read_at_and_summary(self):
        recorder: list[dict[str, Any]] = []
        stubs = _make_stubs(
            summary="The probe document — 2 lines of utf-8.",
            stamp_recorder=recorder,
        )
        result = asyncio.run(_run_workflow(stubs))
        d = _to_dict(result)
        assert d["alfred_read_at"] is not None
        assert d["alfred_read_at"] > 0
        assert "probe document" in d["summary"]
        assert d["extraction_error"] is None
        # The stamp was called exactly once with the success shape.
        assert len(recorder) == 1
        assert recorder[0]["extraction_error"] is None
        assert recorder[0]["summary"] == d["summary"]


class TestUnsupportedMimeShortCircuit:
    def test_stamps_unsupported_mime_when_extract_raises(self):
        recorder: list[dict[str, Any]] = []
        stubs = _make_stubs(
            extract_raises=ExtractionError("unsupported_mime", "test mime"),
            stamp_recorder=recorder,
        )
        result = asyncio.run(_run_workflow(stubs))
        d = _to_dict(result)
        assert d["alfred_read_at"] is None
        assert d["summary"] is None
        assert d["extraction_error"] == "unsupported_mime"
        # The stamp recorded the error code.
        assert recorder[0]["extraction_error"] == "unsupported_mime"


class TestEmptyExtractionShortCircuit:
    def test_empty_text_routes_to_empty_extraction(self):
        recorder: list[dict[str, Any]] = []
        stubs = _make_stubs(
            extract_payload={
                "text": "   ",  # whitespace only
                "char_count": 0,
                "extractor": "text",
                "truncated_for_summary": False,
            },
            stamp_recorder=recorder,
        )
        result = asyncio.run(_run_workflow(stubs))
        d = _to_dict(result)
        assert d["extraction_error"] == "empty_extraction"
        assert d["alfred_read_at"] is None
        assert recorder[0]["extraction_error"] == "empty_extraction"


class TestSummariserFailure:
    def test_summariser_failure_stamps_error(self):
        recorder: list[dict[str, Any]] = []
        stubs = _make_stubs(
            summary_raises=ExtractionError("summariser_failed", "boom"),
            stamp_recorder=recorder,
        )
        result = asyncio.run(_run_workflow(stubs))
        d = _to_dict(result)
        assert d["extraction_error"] == "summariser_failed"
        assert d["alfred_read_at"] is None
        assert recorder[0]["extraction_error"] == "summariser_failed"


class TestStampFailureIsRecorded:
    def test_stamp_failure_still_returns_result(self):
        # The success stamp itself raises — the workflow logs it and
        # returns the in-memory result with extraction_error="stamp_failed".
        recorder: list[dict[str, Any]] = []
        stubs = _make_stubs(
            stamp_raises=RuntimeError("ctrl-api went away"),
            stamp_recorder=recorder,
        )
        result = asyncio.run(_run_workflow(stubs))
        d = _to_dict(result)
        assert d["extraction_error"] == "stamp_failed"
        # And the original stamp attempt was recorded with the success
        # payload (proving we tried).
        assert recorder[0]["alfred_read_at"] is not None
